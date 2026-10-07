#!/usr/bin/env python3
"""Excel de la deuda actual de Rootflow para la asesoría de ENISA.

Fuente: ERP de Rootflow (Supabase) a 07/10/2026 — tablas gastos (forma de pago y reparto entre socios),
aportaciones_socios, reembolsos_socios (vacía) y facturas de gasto pendientes de pago. La contabilidad
automática del ERP (asientos) NO se usa: tiene saldos erróneos.
Uso: python3 scripts/enisa_deuda.py finanzas/Rootflow_ENISA_Deuda.xlsx
"""
import re, sys
from datetime import date as D
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1] if len(sys.argv) > 1 else "finanzas/Rootflow_ENISA_Deuda.xlsx"
CORTE = "07/10/2026"
A = "Arial"
BASE = Font(name=A, size=10); BOLD = Font(name=A, size=10, bold=True)
BLUE = Font(name=A, size=10, color="0000FF"); GREEN = Font(name=A, size=10, color="008000")
GREEN_B = Font(name=A, size=10, color="008000", bold=True)
WHITE = Font(name=A, size=10, bold=True, color="FFFFFF"); TITLE = Font(name=A, size=14, bold=True, color="1D4F37")
SUB = Font(name=A, size=9, italic=True, color="666666"); SEC = Font(name=A, size=11, bold=True, color="1D4F37")
HEAD = PatternFill("solid", fgColor="1D4F37"); TOT = PatternFill("solid", fgColor="E6F2EA")
YEL = PatternFill("solid", fgColor="FFFF00")
LINE = Border(top=Side(style="thin", color="1D4F37"))
EUR = '#,##0.00" €";(#,##0.00" €");"-"'
DATE = "dd/mm/yyyy"
SOCIOS = {"socio_peri": "Pedro Domecq", "socio_nico": "Nicolás Bustamante", "socio_guzman": "Domingo de Guzmán"}
ORDEN = ["socio_peri", "socio_nico", "socio_guzman"]

wb = Workbook()


def hoja(ws, titulo, sub, anchos):
    ws["A1"] = titulo; ws["A1"].font = TITLE
    ws["A2"] = sub; ws["A2"].font = SUB
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[CL(i)].width = w
    ws.sheet_view.showGridLines = False


def cabecera(ws, r, textos):
    for i, t in enumerate(textos, 1):
        c = ws.cell(row=r, column=i, value=t); c.font = WHITE; c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 30


def celda(ws, r, col, v, fmt=None, font=BASE, wrap=False):
    c = ws.cell(row=r, column=col, value=v)
    if fmt: c.number_format = fmt
    c.font = font; c.alignment = Alignment(wrap_text=wrap, vertical="top")
    return c


# ------------------------------------------------------------------ gastos pagados por los socios (ERP)
# (id, fecha, concepto ERP, importe pagado con IVA, pagado por, reparto {socio: importe} o None)
R3 = lambda a, b, c: {"socio_peri": a, "socio_nico": b, "socio_guzman": c}
GS = [
    (18, D(2025, 12, 30), "IONOS WEB 1 (persona)", 1.21, "socio_peri", None),
    (19, D(2026, 1, 30), "WEB IONOS 2 (persona)", 1.21, "socio_peri", None),
    (20, D(2026, 3, 1), "IONOS WEB 3 (persona)", 1.21, "socio_peri", None),
    (27, D(2026, 3, 6), "Alquiler mes de abril calle nueva 16 puerta 6", 535.50, "socio_peri", R3(178.5, 178.5, 178.5)),
    (50, D(2026, 3, 14), "medidor temperatura y humedad (persona)", 21.99, "socio_guzman", None),
    (17, D(2026, 3, 15), "Bandejas leroy merlin (persona)", 34.86, "socio_guzman", None),
    (72, D(2026, 3, 19), "Factura 1 Notaria", 181.50, "socio_peri", R3(60.5, 60.5, 60.5)),
    (73, D(2026, 3, 19), "Factura 2 notaria", 504.40, "socio_peri", R3(168.13, 168.13, 168.14)),
    (74, D(2026, 3, 19), "factura 3 notaria", 98.75, "socio_peri", R3(32.91, 32.91, 32.93)),
    (75, D(2026, 3, 19), "factura 4 notaria", 60.83, "socio_peri", R3(20.27, 20.27, 20.29)),
    (76, D(2026, 3, 19), "factura 5 notaria", 65.34, "socio_peri", R3(21.78, 21.78, 21.78)),
    (49, D(2026, 3, 27), "cables ferretería ( persona) ticket", 16.30, "socio_guzman", None),
    (48, D(2026, 3, 30), "super glue ferreteria ( persona) ticket", 15.60, "socio_guzman", None),
    (24, D(2026, 4, 7), "Suscripcion Claude (persona)", 18.00, "socio_peri", None),
    (25, D(2026, 4, 7), "Tarjetas de visita", 104.96, "socio_peri", None),
    (26, D(2026, 4, 8), "claude (persona)", 5.00, "socio_peri", None),
    (28, D(2026, 4, 9), "claude II (persona)", 5.00, "socio_peri", None),
    (31, D(2026, 4, 9), "Mesa, sillas, enchufes", 144.95, "socio_nico", None),
    (32, D(2026, 4, 9), "Productos Limpieza F Simplificada (persona)", 18.93, "socio_nico", None),
    (33, D(2026, 4, 9), "Candado y Utensilios siembra (persona)", 103.88, "socio_nico", None),
    (34, D(2026, 4, 9), "Temporizador enchufe (persona)", 30.48, "socio_nico", None),
    (35, D(2026, 4, 9), "Inkbird (persona)", 42.48, "socio_nico", None),
    (36, D(2026, 4, 9), "Calentador Cable (persona)", 20.75, "socio_nico", None),
    (37, D(2026, 4, 9), "Plastico Invernadero (persona)", 16.06, "socio_nico", None),
    (38, D(2026, 4, 9), "Manguera (persona)", 20.60, "socio_nico", None),
    (39, D(2026, 4, 9), "Pistola Manguera (persona)", 10.00, "socio_nico", None),
    (40, D(2026, 4, 9), "CUBO NEGRO Wallapop sin factura", 16.58, "socio_nico", None),
    (41, D(2026, 4, 9), "Luces Led (Aliexpress sin factura)", 60.38, "socio_nico", None),
    (42, D(2026, 4, 9), "Luces Led (Aliexpress sin factura)", 27.79, "socio_nico", None),
    (43, D(2026, 4, 9), "Bascula con Bol metalico (sin factura)", 20.79, "socio_nico", None),
    (44, D(2026, 4, 10), "Tarrinas 150cc 250cc (persona)", 142.46, "socio_nico", None),
    (45, D(2026, 4, 14), "Antropic desarrollo erp", 5.00, "socio_peri", None),
    (46, D(2026, 4, 14), "Antropic desarrollo erp 2", 5.00, "socio_peri", None),
    (47, D(2026, 4, 19), "Bandejas de inundado wallapop (sin factura)", 40.00, "socio_guzman", None),
    (51, D(2026, 4, 20), "crocks local (sin factura)", 14.89, "socio_guzman", None),
    (52, D(2026, 4, 20), "crocks 2 (sin factura)", 4.40, "socio_guzman", None),
    (53, D(2026, 4, 21), "crocks 3 (sin factura)", 4.90, "socio_guzman", None),
    (63, D(2026, 4, 21), "Antropic desarrollo erp 3", 5.00, "socio_peri", None),
    (54, D(2026, 4, 24), "Estanteria, sustrato", 213.95, "socio_nico", None),
    (55, D(2026, 4, 24), "tuberias estanteria", 8.68, "socio_guzman", None),
    (56, D(2026, 4, 24), "Tubos estanteria", 82.74, "socio_guzman", None),
    (57, D(2026, 4, 24), "Tijera cortatubos y bomba agua", 37.39, "socio_nico", None),
    (59, D(2026, 4, 25), "sustrato", 14.99, "socio_guzman", None),
    (60, D(2026, 4, 25), "Luces Led (Aliexpress sin factura)", 113.19, "socio_nico", None),
    (61, D(2026, 4, 25), "Impresora etiquetas", 99.95, "socio_nico", None),
    (64, D(2026, 4, 27), "mensual claude pro mayo", 18.00, "socio_peri", None),
    (65, D(2026, 4, 30), "Antropic desarrollo erp 4", 5.00, "socio_peri", None),
    (134, D(2026, 5, 3), "Cotizacion 005 Autonomos", 88.56, "socio_nico", None),
    (62, D(2026, 5, 4), "Alquiler Mayo", 547.61, "socio_peri", R3(182.53, 182.53, 182.55)),
    (66, D(2026, 5, 4), "Antropic desarrollo erp 5", 5.00, "socio_peri", None),
    (69, D(2026, 5, 4), "MD Labels Etiquetas térmicas", 7.41, "socio_peri", None),
    (16, D(2026, 5, 5), "Honorarios alquiler inmobiliaria", 635.25, "socio_peri", R3(211.75, 211.75, 211.75)),
    (67, D(2026, 5, 5), "Antropic desarrollo erp 6", 5.00, "socio_peri", None),
    (68, D(2026, 5, 5), "Antropic desarrollo erp 7", 5.00, "socio_peri", None),
    (71, D(2026, 5, 6), "Antropic desarrollo erp 8", 5.00, "socio_peri", None),
    (77, D(2026, 5, 7), "Cable USB impresora", 4.26, "socio_peri", None),
    (78, D(2026, 5, 7), "Cartucho impresora", 10.50, "socio_peri", None),
    (79, D(2026, 5, 7), "Cables conectores luces Led (aliexpress)", 127.36, "socio_nico", None),
    (81, D(2026, 5, 13), "Cajas verdes plegables", 126.28, "socio_nico", None),
    (85, D(2026, 5, 13), "Antropic desarrollo erp 9", 5.00, "socio_peri", None),
    (105, D(2026, 5, 13), "Registro de la propiedad - certificado", 34.18, "socio_peri", None),
    (82, D(2026, 5, 14), "Xiaomi Redmi + Funda", 124.48, "socio_nico", None),
    (84, D(2026, 5, 14), "Maquina cortadora Posenpro WALLAPOP", 25.22, "socio_nico", None),
    (83, D(2026, 5, 15), "Muesta envases", 5.99, "socio_nico", None),
    (86, D(2026, 5, 18), "Bidon azul wallapop", 25.00, "socio_guzman", None),
    (88, D(2026, 5, 19), "Estanterias y tubos", 274.00, "socio_nico", None),
    (89, D(2026, 5, 19), "buzon y empalme", 12.59, "socio_nico", None),
    (91, D(2026, 5, 20), "Cuba wallapop", 50.00, "socio_nico", None),
    (93, D(2026, 5, 20), "Nevera wallapop", 40.31, "socio_nico", None),
    (108, D(2026, 6, 1), "Conectores luces LED", 78.10, "socio_nico", None),
    (137, D(2026, 6, 3), "Cotizacion 005 Autonomos", 88.56, "socio_nico", None),
    (138, D(2026, 7, 3), "Cotizacion 005 Autonomos", 88.56, "socio_nico", None),
    (139, D(2026, 8, 3), "Cotizacion 005 Autonomos", 88.56, "socio_nico", None),
    (140, D(2026, 9, 3), "Bomba y tuvos", 93.29, "socio_guzman", None),
    (136, D(2026, 9, 4), "Cotizacion 005 Autonomos", 2.95, "socio_nico", None),
    (161, D(2026, 9, 30), "Cotizacion Autonomos", 88.56, "socio_nico", None),
    (160, D(2026, 10, 4), "Gasolina reparto", 221.00, "socio_nico", None),
]
TOT_ERP = {"socio_guzman": 377.64, "socio_nico": 2671.69, "socio_peri": 2885.12}   # sumas del ERP por forma de pago
for s, t in TOT_ERP.items():
    assert abs(sum(g[3] for g in GS if g[4] == s) - t) < 0.005, (s, sum(g[3] for g in GS if g[4] == s))
for g in GS:
    if g[5]: assert abs(sum(g[5].values()) - g[3]) < 0.005, g


def limpiar(c):
    """Concepto presentable: sin anotaciones internas ni nombres de terceros."""
    low = c.lower()
    reglas = [("cotizacion", "Cuota de autónomo del socio"), ("ionos", "IONOS (web y dominio)"),
              ("antropic", "Anthropic (Claude), desarrollo del ERP"), ("claude", "Anthropic (Claude)"),
              ("notaria", "Notaría (constitución de la sociedad)"), ("crocks", "Calzado de trabajo"),
              ("gasolina", "Combustible del reparto a clientes"), ("alquiler mes de abril", "Alquiler del local, abril"),
              ("alquiler mayo", "Alquiler del local, mayo"), ("honorarios alquiler", "Honorarios de la inmobiliaria (alquiler del local)"),
              ("cubo negro", "Cubo"), ("bascula", "Báscula"), ("muesta", "Muestra de envases"), ("bomba y tuvos", "Bomba y tubos"),
              ("maquina cortadora", "Máquina cortadora"), ("bidon", "Bidón"), ("cuba wallapop", "Cuba"),
              ("nevera", "Nevera"), ("bandejas de inundado", "Bandejas de inundación"), ("xiaomi", "Teléfono móvil de empresa"),
              ("luces led", "Luces LED"), ("cables conectores", "Cables y conectores de luces LED"),
              ("productos limpieza", "Productos de limpieza"), ("super glue", "Pegamento (ferretería)"),
              ("cables ferret", "Cables (ferretería)"), ("medidor", "Medidor de temperatura y humedad"),
              ("candado", "Candado y utensilios de siembra"), ("plastico", "Plástico de invernadero")]
    for k, v in reglas:
        if k in low:
            return v
    c = re.sub(r"\s*\(.*?\)", "", c).strip()
    return c[:1].upper() + c[1:]


def segunda_mano(c):
    low = c.lower()
    return "Sin factura (compra a particular o plataforma)" if ("sin factura" in low or "wallapop" in low) else "Factura o ticket"


def tipo(c):
    low = c.lower()
    if "cotizacion" in low: return "Cuota de autónomo del socio"
    if any(k in low for k in ("notaria", "alquiler", "registro", "inmobiliaria")): return "Constitución y local"
    if any(k in low for k in ("antropic", "claude", "ionos")): return "Software y web"
    if "gasolina" in low: return "Reparto"
    if any(k in low for k in ("tarrinas", "etiquetas", "envases", "muesta", "sustrato", "md labels")): return "Consumibles de producción"
    return "Equipamiento y material"


# ================================================================== HOJA: deuda con socios (detalle)
wd = wb.active; wd.title = "Deuda socios (detalle)"
hoja(wd, "Rootflow Hydroponics — Gastos de la sociedad pagados por los socios",
     f"Pendientes de reembolso a {CORTE}: el ERP no registra ningún reembolso. Importes pagados, IVA incluido. "
     "Los gastos de constitución y alquiler que se repartieron a partes iguales figuran a nombre de los tres socios.",
     [12, 46, 26, 22, 34, 13, 15, 15, 15])
cabecera(wd, 4, ["Fecha", "Concepto", "Tipo", "Pagado por", "Justificante", "Importe",
                 SOCIOS["socio_peri"], SOCIOS["socio_nico"], SOCIOS["socio_guzman"]])
d0 = 5
for k, (gid, fe, conc, imp, pag, rep) in enumerate(sorted(GS, key=lambda g: (g[1], g[0]))):
    r = d0 + k
    celda(wd, r, 1, fe, DATE, BLUE); celda(wd, r, 2, limpiar(conc)); celda(wd, r, 3, tipo(conc))
    celda(wd, r, 4, SOCIOS[pag] + (" (repartido entre los tres)" if rep else ""), wrap=False)
    celda(wd, r, 5, segunda_mano(conc)); celda(wd, r, 6, imp, EUR, BLUE)
    for j, s in enumerate(ORDEN):
        if rep:
            celda(wd, r, 7 + j, rep[s], EUR, BLUE)
        else:
            celda(wd, r, 7 + j, f"=IF($D{r}=\"{SOCIOS[s]}\",$F{r},0)", EUR, BASE)
dN = d0 + len(GS) - 1
dt = dN + 1
celda(wd, dt, 2, "TOTAL PENDIENTE DE REEMBOLSO", font=BOLD)
for col in "FGHI":
    c = wd[f"{col}{dt}"]; c.value = f"=SUM({col}{d0}:{col}{dN})"; c.number_format = EUR; c.font = BOLD; c.fill = TOT; c.border = LINE
celda(wd, dt + 1, 2, "Comprobación: reparto = total (debe ser 0)")
c = wd[f"F{dt+1}"]; c.value = f"=ROUND(F{dt}-G{dt}-H{dt}-I{dt},2)"; c.number_format = EUR
wd.freeze_panes = "C5"; wd.auto_filter.ref = f"A4:I{dN}"
DS_TOT = f"'Deuda socios (detalle)'!$F${dt}"
DS_SOC = {s: f"'Deuda socios (detalle)'!${'GHI'[j]}${dt}" for j, s in enumerate(ORDEN)}
DS_RNG_T = f"'Deuda socios (detalle)'!$C${d0}:$C${dN}"
DS_RNG_F = f"'Deuda socios (detalle)'!$F${d0}:$F${dN}"

# ================================================================== HOJA: aportaciones
wa = wb.create_sheet("Aportaciones socios")
hoja(wa, "Rootflow Hydroponics — Aportaciones en efectivo de los socios",
     f"Transferencias de los socios a la cuenta de la sociedad, según el ERP a {CORTE}. Su tratamiento (fondos propios o préstamo) "
     "se elige en la hoja Resumen.", [12, 24, 16, 46])
cabecera(wa, 4, ["Fecha", "Socio", "Importe", "Concepto (ERP)"])
AP = [(D(2026, 5, 12), "socio_guzman", 500, ""), (D(2026, 5, 12), "socio_peri", 500, ""), (D(2026, 5, 18), "socio_nico", 500, ""),
      (D(2026, 5, 29), "socio_nico", 500, ""), (D(2026, 5, 29), "socio_peri", 500, ""), (D(2026, 6, 9), "socio_guzman", 500, ""),
      (D(2026, 7, 3), "socio_peri", 400, "Para pagar el alquiler"), (D(2026, 7, 3), "socio_nico", 400, ""),
      (D(2026, 7, 17), "socio_guzman", 150, ""), (D(2026, 7, 17), "socio_peri", 150, "Para pagar impuestos"),
      (D(2026, 8, 3), "socio_peri", 300, ""), (D(2026, 8, 3), "socio_guzman", 150, ""), (D(2026, 8, 9), "socio_guzman", 150, ""),
      (D(2026, 8, 31), "socio_peri", 426, "Pago del alquiler de septiembre"), (D(2026, 8, 31), "socio_guzman", 400, ""),
      (D(2026, 10, 1), "socio_guzman", 400, ""), (D(2026, 10, 5), "socio_guzman", 100, ""),
      (D(2026, 10, 5), "socio_peri", 500, "Aportación de socios a fondos propios")]
a0 = 5
for k, (fe, s, imp, conc) in enumerate(AP):
    r = a0 + k
    celda(wa, r, 1, fe, DATE, BLUE); celda(wa, r, 2, SOCIOS[s]); celda(wa, r, 3, imp, EUR, BLUE); celda(wa, r, 4, conc)
aN = a0 + len(AP) - 1
celda(wa, aN + 1, 2, "TOTAL", font=BOLD)
c = wa[f"C{aN+1}"]; c.value = f"=SUM(C{a0}:C{aN})"; c.number_format = EUR; c.font = BOLD; c.fill = TOT; c.border = LINE
for j, s in enumerate(ORDEN):
    r = aN + 3 + j
    celda(wa, r, 2, SOCIOS[s]); c = wa[f"C{r}"]
    c.value = f'=SUMIFS($C${a0}:$C${aN},$B${a0}:$B${aN},B{r})'; c.number_format = EUR
AP_TOT = f"'Aportaciones socios'!$C${aN+1}"
AP_SOC = {s: f"'Aportaciones socios'!$C${aN+3+j}" for j, s in enumerate(ORDEN)}

# ================================================================== HOJA: acreedores y AAPP
wc = wb.create_sheet("Proveedores y Hacienda")
hoja(wc, "Rootflow Hydroponics — Facturas pendientes de pago y obligaciones con Hacienda",
     f"Según el ERP a {CORTE}. Las retenciones son una estimación a partir de las facturas del arrendador: confirmar con la gestoría.",
     [12, 30, 52, 15, 16, 44])
cabecera(wc, 4, ["Fecha", "Acreedor", "Concepto", "Importe", "Vencimiento", "Nota"])
PEND = [(D(2026, 9, 30), "Bellbro Renting and Development, S.L.", "Alquiler del local, octubre 2026", 535.50, D(2026, 10, 5),
         "525,00 € + IVA − retención IRPF"),
        (D(2026, 9, 30), "Bellbro Renting and Development, S.L.", "Luz repercutida, 15/07 a 15/09/2026", 137.51, D(2026, 10, 5),
         "134,82 € + IVA − retención IRPF"),
        (D(2026, 9, 30), "Bellbro Renting and Development, S.L.", "Agua, 10/07 a 08/09/2026", 30.71, D(2026, 10, 5), "")]
for k, (fe, acr, conc, imp, ven, nota) in enumerate(PEND):
    r = 5 + k
    celda(wc, r, 1, fe, DATE, BLUE); celda(wc, r, 2, acr); celda(wc, r, 3, conc); celda(wc, r, 4, imp, EUR, BLUE)
    celda(wc, r, 5, ven, DATE, BLUE); celda(wc, r, 6, nota)
pN = 5 + len(PEND) - 1
celda(wc, pN + 1, 3, "Total facturas de proveedores pendientes de pago", font=BOLD)
c = wc[f"D{pN+1}"]; c.value = f"=SUM(D5:D{pN})"; c.number_format = EUR; c.font = BOLD; c.fill = TOT; c.border = LINE
PROV_TOT = f"'Proveedores y Hacienda'!$D${pN+1}"
celda(wc, pN + 2, 6, "Vencimiento: primeros días del mes (pago domiciliado habitual).", font=SUB)

h0 = pN + 4
celda(wc, h0, 1, "HACIENDA — retenciones de IRPF del alquiler (modelo 115) pendientes de ingresar", font=BOLD)
cabecera(wc, h0 + 1, ["Fecha factura", "Arrendador", "Factura", "Retención 19 %", "A ingresar hasta", "Nota"])
# El 115 se devenga al pagar (o ser exigible) la renta: agosto y septiembre se pagaron en el 3T;
# las facturas del 30/09 están sin pagar y su retención irá al 4T.
N3T = "3T: alquiler pagado en el trimestre."
N4T = "4T: factura del 30/09 aún sin pagar; la retención se ingresa al pagarla."
RET = [(D(2026, 7, 28), "Alquiler agosto", 99.75, D(2026, 10, 20), N3T),
       (D(2026, 8, 31), "Alquiler septiembre", 99.75, D(2026, 10, 20), N3T),
       (D(2026, 9, 30), "Alquiler octubre", 99.75, D(2027, 1, 20), N4T),
       (D(2026, 9, 30), "Luz repercutida 15/07–15/09", 25.62, D(2027, 1, 20), N4T)]
for k, (fe, conc, ret, ven, nota) in enumerate(RET):
    r = h0 + 2 + k
    celda(wc, r, 1, fe, DATE, BLUE); celda(wc, r, 2, "Bellbro Renting and Development, S.L."); celda(wc, r, 3, conc)
    celda(wc, r, 4, ret, EUR, BLUE); celda(wc, r, 5, ven, DATE, BLUE)
    celda(wc, r, 6, nota)
hN = h0 + 2 + len(RET) - 1
celda(wc, hN + 1, 3, "Total retenciones pendientes de ingresar (estimado)", font=BOLD)
c = wc[f"D{hN+1}"]; c.value = f"=SUM(D{h0+2}:D{hN})"; c.number_format = EUR; c.font = BOLD; c.fill = TOT; c.border = LINE
RET_TOT = f"'Proveedores y Hacienda'!$D${hN+1}"
celda(wc, hN + 3, 1, "IVA: en 2026 el IVA soportado (21 % en compras y alquiler) supera al repercutido (4 % en ventas), "
      "así que la sociedad tiene saldo a compensar o devolver, no deuda. Impuesto sobre Sociedades: sin cuota (pérdidas). "
      "Seguridad Social: la sociedad no tiene empleados.", font=SUB)

# ================================================================== RESUMEN
ws = wb.create_sheet("Resumen", 0)
ws.sheet_view.showGridLines = False
for i, w in enumerate([62, 18, 74], 1):
    ws.column_dimensions[CL(i)].width = w
ws["A1"] = "Rootflow Hydroponics, S.L. — Deuda actual"; ws["A1"].font = TITLE
ws["A2"] = f"Situación a {CORTE} según los registros de la empresa (ERP). Importes en euros."; ws["A2"].font = SUB


def lin(r, et, f, font=None, nota="", fmt=EUR, fill=None):
    ws[f"A{r}"] = et; ws[f"A{r}"].font = font or BASE
    c = ws[f"B{r}"]; c.value = f; c.number_format = fmt; c.font = font or GREEN
    if fill: c.fill = fill; ws[f"A{r}"].fill = fill
    if nota: ws[f"C{r}"] = nota; ws[f"C{r}"].font = SUB; ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")


ws["A4"] = "¿Las aportaciones en efectivo de los socios son un préstamo a la sociedad?"; ws["A4"].font = BOLD
c = ws["B4"]; c.value = "No"; c.font = BLUE; c.fill = YEL; c.alignment = Alignment(horizontal="center")
ws["C4"] = ("«No» = aportaciones a fondos propios (cuenta 118), que no son deuda. «Sí» = préstamo de socios. "
            "Confirmar con la gestoría cómo están contabilizadas."); ws["C4"].font = SUB
ws["C4"].alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False); ws.add_data_validation(dv); dv.add("B4")
ws.row_dimensions[4].height = 30

ws["A6"] = "1 · DEUDA FINANCIERA CON ENTIDADES DE CRÉDITO"; ws["A6"].font = SEC
lin(7, "Préstamos, pólizas de crédito, leasing o renting financiero", 0, font=None,
    nota="No hay ninguno registrado en el ERP.")
ws["B7"].font = BLUE

ws["A9"] = "2 · DEUDA CON SOCIOS"; ws["A9"].font = SEC
lin(10, "Gastos de la sociedad pagados por los socios, pendientes de reembolso", f"={DS_TOT}",
    nota="Detalle en la hoja «Deuda socios (detalle)». Sin intereses ni vencimiento pactado.")
for j, s in enumerate(ORDEN):
    lin(11 + j, f"   · {SOCIOS[s]}", f"={DS_SOC[s]}")
lin(14, "   · de los cuales, cuotas de autónomo del socio asumidas por la sociedad",
    f'=SUMIFS({DS_RNG_F},{DS_RNG_T},"Cuota de autónomo del socio")', nota="Confirmar con la gestoría su tratamiento.")
lin(15, "Aportaciones en efectivo de los socios (si se tratan como préstamo)", f'=IF($B$4="Sí",{AP_TOT},0)',
    nota="Con «No» en B4 cuentan como fondos propios y no suman aquí.")
lin(16, "TOTAL DEUDA CON SOCIOS", "=B10+B15", font=BOLD, fill=TOT)

ws["A18"] = "3 · PROVEEDORES Y ACREEDORES COMERCIALES"; ws["A18"].font = SEC
lin(19, "Facturas recibidas pendientes de pago", f"={PROV_TOT}", nota="Alquiler de octubre, luz y agua del arrendador (30/09).")

ws["A21"] = "4 · ADMINISTRACIONES PÚBLICAS"; ws["A21"].font = SEC
lin(22, "Retenciones de IRPF del alquiler pendientes de ingresar (modelo 115, estimado)", f"={RET_TOT}",
    nota="199,50 € del 3T, a ingresar hasta el 20/10/2026; el resto, de las facturas del 30/09, va al 4T. "
         "IVA con saldo a favor; sin deuda con la Seguridad Social.")
ws.row_dimensions[22].height = 30

lin(24, "DEUDA TOTAL", "=B7+B16+B19+B22", font=Font(name=A, size=11, bold=True), fill=TOT)
lin(25, "   · de la cual, con socios", "=B16")
lin(26, "   · de la cual, con terceros", "=B7+B19+B22")

ws["A28"] = "PARA REFERENCIA — FINANCIACIÓN APORTADA POR LOS SOCIOS"; ws["A28"].font = SEC
lin(29, "Aportaciones en efectivo a la cuenta de la sociedad", f"={AP_TOT}", nota="Hoja «Aportaciones socios».")
for j, s in enumerate(ORDEN):
    lin(30 + j, f"   · {SOCIOS[s]}", f"={AP_SOC[s]}")
lin(33, "Gastos pagados directamente por los socios", "=B10")
lin(34, "TOTAL FINANCIADO POR LOS SOCIOS", "=B29+B33", font=BOLD, fill=TOT,
    nota="Entre la deuda con socios y las aportaciones, los socios han financiado toda la actividad hasta hoy.")
ws["A36"] = ("Hojas: Deuda socios (detalle) · Aportaciones socios · Proveedores y Hacienda. "
             "No incluye el capital social escriturado."); ws["A36"].font = SUB

for w in wb.worksheets:
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0
wb.save(OUT)
print("escrito", OUT)
