#!/usr/bin/env python3
"""Extrae de un modelo recalculado la tabla de cifras que alimentan los decks."""
import openpyxl, sys, json

F_KG, F_ING, F_PLANT, F_OCUP, F_POS = 26, 36, 31, 25, 24
F_EBITDA_A, F_DSCR_A = 142, 144

def cifras(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    h, pr, cx = wb["Hipótesis"], wb["Producto"], wb["CAPEX"]
    o = {}
    o["kg_band_mes"]     = h["F23"].value
    o["ciclos_mes"]      = h["F24"].value
    o["precio_medio_kg"] = pr["D35"].value
    o["margen_bruto_pf"] = pr["D36"].value
    o["ingreso_band_mes"]= pr["T32"].value
    o["dias_muertos"]    = pr["D7"].value
    o["bandejas_rack"]   = h["D13"].value
    o["racks"]           = h["F26"].value
    o["posiciones_nave"] = (h["D13"].value or 0) * (h["F26"].value or 0)
    o["m2_nave"]         = (h["D14"].value or 0) * (h["F26"].value or 0)
    o["m2_bandeja"]      = h["D11"].value
    o["band_rack_antiguo"] = h["D8"].value

    # rendimiento por variedad activa (para el pico en monocultivo)
    act = []
    for r in range(16, 32):
        if (pr.cell(row=r, column=3).value or 0) == 1:
            act.append(dict(var=pr.cell(row=r, column=1).value,
                            dias_luz=pr.cell(row=r, column=8).value,
                            ciclos=pr.cell(row=r, column=10).value,
                            g=pr.cell(row=r, column=11).value,
                            kg_band_mes=pr.cell(row=r, column=13).value))
    o["activas"] = act
    rab = [a for a in act if "ábano" in str(a["var"])]
    o["kg_band_mes_rabano"] = max((a["kg_band_mes"] for a in rab), default=0)

    # CAPEX (coordenadas fijas)
    o["capex"] = {
        "adaptacion_nave": cx["C5"].value, "automatizacion": cx["C10"].value,
        "racks": cx["C13"].value, "instalacion": cx["C14"].value,
        "vehiculo": cx["C15"].value, "mobiliario": cx["C16"].value,
        "licencias": cx["C17"].value, "legal": cx["C18"].value,
        "subtotal": cx["C19"].value, "contingencia": cx["C21"].value,
        "capex_total": cx["C22"].value, "colchon_circulante": cx["C25"].value,
        "bloque_obra_bottomup": cx["C35"].value if cx["C35"].value is not None else None,
        "bloque_obra_aplicado": cx["C38"].value if cx["C38"].value is not None else None,
    }
    o["necesidad"]  = cx["C26"].value
    o["ronda_total"]= (h["D93"].value or 0)+(h["D105"].value or 0)+(h["D113"].value or 0)
    o["cobertura"]  = o["ronda_total"]/o["necesidad"] if o["necesidad"] else None
    o["pct_inversor"]     = h["D119"].value
    o["importe_inversor"] = h["D92"].value
    o["importe_enisa"]    = h["D104"].value

    esc = {}
    for hoja, nom in (("Esc_Conservador","conservador"),("Esc_Base","base"),("Esc_Optimista","optimista")):
        b = wb[hoja]; g = lambda r: b.cell(row=r, column=3).value
        d = dict(tir=g(128), moic=g(129), equity_salida=g(157), cobro=g(158),
                 dscr_min=g(145), caja_min=g(146), cumple=g(165),
                 kg_mes_regimen=b.cell(row=F_KG, column=62).value,
                 ocupacion_regimen=b.cell(row=F_OCUP, column=62).value,
                 posiciones_regimen=b.cell(row=F_POS, column=62).value)
        d["anual"] = []
        for i, col in enumerate(range(3, 10), 1):
            c0 = 3+(i-1)*12
            ing  = sum((b.cell(row=F_ING, column=c).value or 0) for c in range(c0, c0+12))
            kg   = sum((b.cell(row=F_KG,  column=c).value or 0) for c in range(c0, c0+12))
            ebit = b.cell(row=F_EBITDA_A, column=col).value or 0
            d["anual"].append(dict(ano=i, ingresos=ing, ebitda=ebit,
                                   margen=(ebit/ing if ing else 0), kg=kg,
                                   plantilla=b.cell(row=F_PLANT, column=c0+11).value,
                                   dscr=b.cell(row=F_DSCR_A, column=col).value))
        esc[nom] = d
    o["escenarios"] = esc
    return o

if __name__ == "__main__":
    print(json.dumps(cifras(sys.argv[1]), ensure_ascii=False, indent=1, default=str))
