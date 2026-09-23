#!/usr/bin/env python3
"""V16 = V15 + revision operativa del socio de produccion.

BLOQUE A (aplicado)
  Rendimientos por bandeja: cilantro 120->145, albahaca 100->90,
  rabano daikon y rose 175->245 (solo los ACTIVOS).
BLOQUE B (parametrizado, sin cambiar el resultado por defecto)
  B1 Dias muertos entre ciclos: input ya existente (Producto!D7), documentado.
  B2 Presupuesto real del bloque nave+racks+acondicionamiento, con % de ajuste.
  B4 Hoja Tarifas para cargar el export del ERP y calcular el precio ponderado.
NO SE TOCA: estructura de capital, selector de escenarios, Mano_de_obra, OPEX, precios.
"""
import openpyxl, shutil, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ORIGEN  = "modelo/Rootflow_Modelo_Financiero_V15_equity.xlsx"
DESTINO = "modelo/Rootflow_Modelo_Financiero_V16.xlsx"
AZUL = PatternFill("solid", fgColor="DCE9F7")     # convencion del modelo: input
AMAR = PatternFill("solid", fgColor="FFF3CD")
THIN = Side(style="thin", color="BFC9C2")
BOR  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# rendimientos nuevos: fila -> (variedad, g/bandeja)
RENDIMIENTOS = {17: ("Cilantro split", 145),
                18: ("Albahaca genovesa", 90),
                19: ("Rábano daikon", 245),
                21: ("Rábano rose", 245)}

def main(destino=DESTINO):
    shutil.copy(ORIGEN, destino)
    wb = openpyxl.load_workbook(destino, data_only=False)
    pr, cx = wb["Producto"], wb["CAPEX"]

    # ---------- BLOQUE A · rendimientos ----------
    for fila, (nombre, g) in RENDIMIENTOS.items():
        pr.cell(row=fila, column=11).value = g        # K = Rendim. (g/band)
        pr.cell(row=fila, column=11).fill  = AZUL
    # nota de trazabilidad
    pr["A38"] = "REVISIÓN OPERATIVA (socio de producción) — 23/09/2026"
    pr["A38"].font = Font(bold=True)
    pr["A39"] = ("Rendimientos actualizados: cilantro 120→145 g · albahaca 100→90 g · "
                 "rábano daikon y rose 175→245 g (solo los ACTIVOS).")
    pr["A40"] = ("PENDIENTE DE CONFIRMAR: 'Rábano red' (INACTIVO) sigue en 200 g y "
                 "'Micromezclum' (mix de 7 variedades) sigue en 170 g. Si el mix lleva rábano, "
                 "su rendimiento sube también — revisar con producción.")
    pr["A41"] = ("OJO: los rábanos ACTIVOS tienen 3 días de luz (10,1 ciclos/mes). "
                 "El 'rábano de ciclo corto' del deck (2 días de luz, 15 ciclos/mes) es el rábano red, que está INACTIVO.")
    for r in (39, 40, 41):
        pr.cell(row=r, column=1).alignment = Alignment(wrap_text=False)

    # ---------- BLOQUE B1 · días muertos entre ciclos ----------
    pr["D7"].fill = AZUL
    pr["F7"] = ("INPUT. Ciclos/mes = 30,42 ÷ (días de luz + días muertos). "
                "Con 1 día muerto, un rábano de 3 días de luz pasa de 10,1 a 7,6 ciclos/mes.")

    # ---------- BLOQUE B2 · presupuesto real del bloque de obra ----------
    cx["A34"] = "PRESUPUESTO DE OBRA — pendiente del presupuesto real"
    cx["A34"].font = Font(bold=True)
    cx["A35"] = "Bloque nave + racks + acondicionamiento — bottom-up actual (€)"
    cx["C35"] = "=C5+C13+C14+C16"
    cx["A36"] = "PRESUPUESTO REAL del bloque (0 = usar el bottom-up de arriba)"
    cx["C36"] = 0; cx["C36"].fill = AZUL
    cx["D36"] = "Cuando llegue el presupuesto, escríbelo aquí y el modelo entero se recalcula."
    cx["A37"] = "Ajuste % sobre el bloque (0 = ninguno · 0,25 = +25 %)"
    cx["C37"] = 0; cx["C37"].fill = AZUL
    cx["D37"] = "Para la sensibilidad rápida sin tocar el desglose."
    cx["A38"] = "BLOQUE DE OBRA APLICADO (€)"
    cx["A38"].font = Font(bold=True)
    cx["C38"] = "=IF($C$36>0,$C$36,$C$35*(1+$C$37))"
    for r in (35, 36, 37, 38):
        for c in ("A", "C", "D"):
            cx[f"{c}{r}"].border = BOR

    # el subtotal pasa a usar el bloque aplicado en vez de las partidas sueltas
    cx["C19"] = "=$C$38+C10+C15+C17+C18"
    cx["E19"] = "=$C$38+E10+C15+C17+C18"
    cx["F19"] = "=$C$38+F10+C15+C17+C18"
    cx["G19"] = "=$C$38+G10+C15+C17+C18"
    cx["D19"] = "Nave+racks+instalación+utillaje vía C38; el resto, partida a partida."

    # ---------- BLOQUE B4 · hoja de tarifas del ERP ----------
    if "Tarifas" in wb.sheetnames:
        del wb["Tarifas"]
    tf = wb.create_sheet("Tarifas", wb.sheetnames.index("Producto") + 1)
    tf["A1"] = "ROOTFLOW — Tarifas reales por cliente, variedad y formato"
    tf["A1"].font = Font(bold=True, size=14, color="1F4D2B")
    tf["A2"] = ("Pega aquí el export del ERP. El precio medio ponderado por variedad se calcula solo "
                "y se copia a mano en Producto!F (Precio DIR €/tarrina). NADA de esta hoja alimenta "
                "el modelo todavía: primero hay que revisar los datos.")
    cab = ["Cliente", "Variedad", "Canal (DIR/DIST)", "Formato (g/tarrina)",
           "Precio €/tarrina", "Tarrinas/mes", "€/mes — calc", "kg/mes — calc", "€/kg — calc"]
    tf.append([]); tf.append(cab)
    for i, t in enumerate(cab, start=1):
        c = tf.cell(row=4, column=i)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="1F4D2B")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        tf.column_dimensions[get_column_letter(i)].width = [22,22,17,17,15,13,14,14,12][i-1]
    for r in range(5, 205):
        tf.cell(row=r, column=7).value = f"=IF(N(E{r})*N(F{r})=0,\"\",E{r}*F{r})"
        tf.cell(row=r, column=8).value = f"=IF(N(D{r})*N(F{r})=0,\"\",D{r}*F{r}/1000)"
        tf.cell(row=r, column=9).value = f"=IFERROR(G{r}/H{r},\"\")"
        for c in range(1, 10):
            tf.cell(row=r, column=c).border = BOR
            if c <= 6: tf.cell(row=r, column=c).fill = AZUL

    tf["K4"] = "RESUMEN POR VARIEDAD (se rellena solo desde la tabla)"
    tf["K4"].font = Font(bold=True)
    tf["K5"] = "Variedad"; tf["L5"] = "€/mes"; tf["M5"] = "kg/mes"; tf["N5"] = "€/kg ponderado"
    tf["O5"] = "€/tarrina medio"
    for i, t in enumerate(("K5","L5","M5","N5","O5"), start=11):
        pass
    for co in ("K5","L5","M5","N5","O5"):
        tf[co].font = Font(bold=True, color="FFFFFF", size=10)
        tf[co].fill = PatternFill("solid", fgColor="1F4D2B")
    # una fila por variedad del portfolio
    for i, fila in enumerate(range(16, 32), start=6):
        tf.cell(row=i, column=11).value = f"=Producto!A{fila}"
        tf.cell(row=i, column=12).value = f'=SUMIF($B$5:$B$204,$K{i},$G$5:$G$204)'
        tf.cell(row=i, column=13).value = f'=SUMIF($B$5:$B$204,$K{i},$H$5:$H$204)'
        tf.cell(row=i, column=14).value = f'=IFERROR(L{i}/M{i},"")'
        tf.cell(row=i, column=15).value = f'=IFERROR(SUMIF($B$5:$B$204,$K{i},$G$5:$G$204)/SUMIF($B$5:$B$204,$K{i},$F$5:$F$204),"")'
        for c in range(11, 16): tf.cell(row=i, column=c).border = BOR
    tf["K23"] = "TOTAL / PONDERADO"; tf["K23"].font = Font(bold=True)
    tf["L23"] = "=SUM(L6:L21)"; tf["M23"] = "=SUM(M6:M21)"; tf["N23"] = '=IFERROR(L23/M23,"")'
    for co in ("L23","M23","N23"): tf[co].font = Font(bold=True)
    tf["K25"] = "Precio medio ponderado que usa HOY el modelo (Producto!D35):"
    tf["N25"] = "=Producto!D35"
    tf["K26"] = "Diferencia vs. tarifas reales (€/kg):"
    tf["N26"] = '=IFERROR(N23-N25,"")'
    for c, w in (("K",26),("L",14),("M",14),("N",16),("O",16)):
        tf.column_dimensions[c].width = w

    wb.save(destino)
    print(f"V16 escrito en {destino}")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DESTINO)
