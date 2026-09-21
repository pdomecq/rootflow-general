#!/usr/bin/env python3
"""Genera el Excel multi-pestana a partir del CSV consolidado."""
import csv, os, re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "output")
SRC = os.path.join(OUT, "inversores_rootflow_consolidado.csv")
DST = os.path.join(OUT, "inversores_rootflow_consolidado.xlsx")

VERDE = "1F6F4A"; VERDE_CLARO = "E8F3EC"
AMBAR = "B7791F"; ROJO = "9B2C2C"; GRIS = "4A5568"
HDR = Font(bold=True, color="FFFFFF", size=10)
FILL_HDR = PatternFill("solid", fgColor=VERDE)
THIN = Side(style="thin", color="D5DDD8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FILL_A = PatternFill("solid", fgColor="D4EDDA")
FILL_B = PatternFill("solid", fgColor="FFF3CD")
FILL_C = PatternFill("solid", fgColor="F1F3F5")

ANCHOS = {"entidad":34,"tipo":17,"subtipo":26,"pais":10,"ciudad":16,"web":30,
          "tesis_resumen":52,"evidencia_agro_food":52,"angulo_pitch":52,
          "red_flags":40,"persona_contacto":24,"cargo":26,"linkedin_url":30,
          "email_publico":26,"fuente_url_1":34,"fuente_url_2":34,"instrumentos":24,
          "siguiente_paso":46,"id":12,"score":7,"tier":6}

def leer():
    with open(SRC, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=";")), None

def siguiente_paso(r):
    """Accion concreta sugerida segun via de entrada y tipo."""
    via, tipo = r.get("via_entrada",""), r.get("tipo","")
    ent = r.get("entidad","")
    if r.get("estado") == "ya_en_pipeline":
        return "YA EN CURSO - retomar y cerrar siguiente hito"
    if via == "evento":
        return "Agendar encuentro en feria - preparar one-pager impreso"
    if via == "intro_calida":
        return "Buscar intro calida en red comun antes de escribir"
    if via == "red_angels":
        return "Aplicar via la red de angels - pitch en foro"
    if via == "formulario_web":
        return "Enviar via formulario web + seguimiento en 5 dias"
    if via == "email_generico" and (r.get("email_publico") or "-") not in ("-",""):
        return f"Email a {r.get('email_publico')} con resumen + deck"
    if tipo in ("industrial_agro","cadena_valor"):
        return "Contacto comercial primero (cliente/marca blanca), luego inversion"
    return "Investigar via de entrada - LinkedIn del responsable"

def estilar(ws, cols, rows, colorear_tier=True):
    ws.append(cols)
    for c in range(1, len(cols)+1):
        cell = ws.cell(row=1, column=c); cell.font = HDR; cell.fill = FILL_HDR
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = ANCHOS.get(cols[c-1], 15)
    ws.row_dimensions[1].height = 28
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    tier_i = cols.index("tier")+1 if "tier" in cols else None
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER; cell.font = Font(size=9)
        if colorear_tier and tier_i:
            t = ws.cell(row=row[0].row, column=tier_i).value
            fill = {"A":FILL_A,"B":FILL_B,"C":FILL_C}.get(t)
            if fill:
                for cell in row: cell.fill = fill
    # score numerico
    if "score" in cols:
        si = cols.index("score")+1
        for rr in range(2, ws.max_row+1):
            v = ws.cell(row=rr, column=si).value
            try: ws.cell(row=rr, column=si).value = int(float(v))
            except Exception: pass
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

def main():
    rows, _ = leer()
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    COLS_FULL = ["entidad","tipo","subtipo","ciudad","pais","web","tesis_resumen",
        "evidencia_agro_food","ticket_min_eur","ticket_max_eur","instrumentos",
        "acepta_minoritario_sin_control","ultima_operacion_fecha","persona_contacto","cargo",
        "linkedin_url","email_publico","via_entrada","angulo_pitch","red_flags","conflicto",
        "score","tier","confianza","estado","fuente_url_1","fuente_url_2","fecha_verificacion","id"]

    # --- Portada -------------------------------------------------------------
    ws = wb.create_sheet("Portada")
    tot = len(rows)
    nA = sum(1 for r in rows if r["tier"]=="A"); nB = sum(1 for r in rows if r["tier"]=="B")
    nC = tot-nA-nB
    nuevos = sum(1 for r in rows if r["estado"]=="nuevo")
    info = [
        ("ROOTFLOW HYDROPONICS — Mapa de inversores para la ronda", ""),
        ("", ""),
        ("Fecha de elaboracion", "21/09/2026"),
        ("Ronda objetivo", "200.000 EUR = 100.000 privado + 100.000 ENISA"),
        ("Necesidad de financiacion (Modelo V14)", "161.278 EUR  ->  cobertura 1,24x"),
        ("Instrumento preferido", "Prestamo participativo + equity kicker ~2,5%"),
        ("Ticket buscado", "50-100k lead  /  20-50k sindicado"),
        ("Retorno inversor (escenario Base)", "TIR 33,0%  ·  MOIC 2,39x  ·  payback 33 meses"),
        ("Hito critico", "Term sheet firmado el 31/10/2026"),
        ("", ""),
        ("Entidades unicas verificadas", tot),
        ("Tier A (prioridad maxima)", nA),
        ("Tier B (segunda oleada)", nB),
        ("Tier C (baja prioridad / descartables)", nC),
        ("Nuevas (no estaban en el pipeline)", nuevos),
        ("", ""),
        ("CONFIDENCIAL — uso interno. No publicar ni enviar a terceros.", ""),
    ]
    for i,(a,b) in enumerate(info, start=1):
        ws.cell(row=i, column=1, value=a); ws.cell(row=i, column=2, value=b)
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 56
    ws["A1"].font = Font(bold=True, size=15, color=VERDE)
    for i in range(3, len(info)+1):
        ws.cell(row=i, column=1).font = Font(bold=True, size=10)
        ws.cell(row=i, column=2).font = Font(size=10)
    ws.cell(row=len(info), column=1).font = Font(bold=True, size=10, color=ROJO)

    # --- Top 30 --------------------------------------------------------------
    activos = [r for r in rows if r["estado"] != "ya_en_pipeline"]
    top = sorted(activos, key=lambda r: -int(float(r["score"] or 0)))[:30]
    for r in top: r["siguiente_paso"] = siguiente_paso(r)
    COLS_TOP = ["entidad","tipo","ciudad","web","score","tier","ticket_min_eur","ticket_max_eur",
                "instrumentos","angulo_pitch","via_entrada","siguiente_paso","persona_contacto",
                "cargo","email_publico","red_flags","estado","fuente_url_1"]
    estilar(wb.create_sheet("Top 30"), COLS_TOP, top)

    # --- Consolidado ---------------------------------------------------------
    estilar(wb.create_sheet("Consolidado"), COLS_FULL, rows)

    # --- Una pestana por segmento -------------------------------------------
    SEG = {
        "Industriales agro": lambda r: r["tipo"] in ("industrial_agro",),
        "Cadena de valor": lambda r: r["tipo"] in ("cadena_valor",),
        "Family offices": lambda r: r["tipo"] in ("family_office",),
        "Angels y redes": lambda r: r["tipo"] in ("angel","red_angels","crowdfunding"),
        "VC CVC e impacto": lambda r: r["tipo"] in ("vc","cvc","impacto","aceleradora"),
        "Deuda alternativa": lambda r: r["tipo"] in ("venture_debt","deuda_alternativa","publico_coinversor","banca"),
    }
    for nombre, pred in SEG.items():
        sub = [r for r in rows if pred(r)]
        if sub: estilar(wb.create_sheet(nombre), COLS_FULL, sub)
    otros = [r for r in rows if not any(p(r) for p in SEG.values())]
    if otros: estilar(wb.create_sheet("Otros"), COLS_FULL, otros)

    # --- Descartados ---------------------------------------------------------
    desc = [r for r in rows if r["tier"]=="C" or r["conflicto"]=="competidor"]
    for r in desc:
        motivo = []
        if r["conflicto"]=="competidor": motivo.append("COMPETIDOR - tratar como posible adquirente, no como inversor")
        if r["tier"]=="C": motivo.append("Encaje bajo (score <55)")
        if (r["red_flags"] or "-") not in ("-",""): motivo.append(r["red_flags"])
        r["motivo_descarte"] = " | ".join(motivo)
    if desc:
        estilar(wb.create_sheet("Descartados"),
                ["entidad","tipo","ciudad","web","score","tier","motivo_descarte","red_flags","fuente_url_1"], desc)

    # --- Fuentes -------------------------------------------------------------
    fuentes = defaultdict(list)
    for r in rows:
        for k in ("fuente_url_1","fuente_url_2"):
            u = (r.get(k) or "").strip()
            if u and u != "-" and u.startswith("http"):
                m = re.sub(r"^https?://(www\.)?", "", u).split("/")[0]
                fuentes[m].append(r["entidad"])
    ws = wb.create_sheet("Fuentes")
    estilar(ws, ["dominio","n_entidades","ejemplos"],
            [{"dominio":d,"n_entidades":len(e),"ejemplos":", ".join(sorted(set(e))[:6])}
             for d,e in sorted(fuentes.items(), key=lambda x:-len(x[1]))], colorear_tier=False)
    ws.column_dimensions["A"].width = 36; ws.column_dimensions["C"].width = 70

    # --- Seguimiento (CRM listo para rellenar) -------------------------------
    ws = wb.create_sheet("Seguimiento")
    cols_seg = ["entidad","tipo","score","tier","persona_contacto","via_entrada",
                "siguiente_paso","fecha_1er_contacto","canal","estado_contacto",
                "NDA","modelo_enviado","visita","proxima_accion","fecha_proxima_accion","notas"]
    filas_seg = []
    for r in top:
        filas_seg.append({
            "entidad": r["entidad"], "tipo": r["tipo"], "score": r["score"], "tier": r["tier"],
            "persona_contacto": r.get("persona_contacto",""), "via_entrada": r.get("via_entrada",""),
            "siguiente_paso": r.get("siguiente_paso",""),
            "fecha_1er_contacto":"", "canal":"", "estado_contacto":"pendiente",
            "NDA":"", "modelo_enviado":"", "visita":"", "proxima_accion":"", 
            "fecha_proxima_accion":"", "notas":"",
        })
    estilar(ws, cols_seg, filas_seg)
    for c,w in (("H",16),("I",12),("J",16),("K",8),("L",16),("M",10),("N",30),("O",18),("P",40)):
        ws.column_dimensions[c].width = w
    # validacion de estado
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list",
        formula1='"pendiente,contactado,interesado,NDA firmado,numeros enviados,visita hecha,oferta,descartado"',
        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{max(ws.max_row,2)}")
    dv2 = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws.add_data_validation(dv2)
    for col in ("K","L","M"):
        dv2.add(f"{col}2:{col}{max(ws.max_row,2)}")

    wb.save(DST)
    print(f"Escrito {DST}")
    print(f"Pestanas: {wb.sheetnames}")
    print(f"Top 30 cabecera: {[r['entidad'] for r in top[:10]]}")

if __name__ == "__main__":
    main()
