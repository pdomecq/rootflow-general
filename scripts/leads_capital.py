#!/usr/bin/env python3
"""Anade al Excel la pestana 'Leads capital': todas las entidades que pueden
aportar CAPITAL (no solo deuda), clasificadas por canal de primer contacto
y preparadas como hoja de seguimiento."""
import csv, os
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC  = os.path.join(BASE, "output", "inversores_rootflow_consolidado.csv")
DST  = os.path.join(BASE, "output", "inversores_rootflow_consolidado.xlsx")

# Tipos que pueden ENTRAR EN CAPITAL. Se excluyen banca, venture debt y deuda
# alternativa: prestan, no invierten.
TIPOS_CAPITAL = {"family_office","angel","red_angels","vc","cvc","impacto",
                 "aceleradora","industrial_agro","cadena_valor"}
VACIO = ("","-","desconocido","desconocida","n/a")

def puede_poner_capital(r):
    if r["tipo"] in TIPOS_CAPITAL:
        return True
    # crowdfunding y coinversores publicos solo si hacen equity
    if r["tipo"] in ("crowdfunding","publico_coinversor"):
        return "equity" in (r.get("instrumentos") or "").lower()
    return False

def canal(r):
    """Devuelve (canal, quien, accion_concreta, dato_de_contacto)."""
    email = (r.get("email_publico") or "").strip()
    via   = r.get("via_entrada","")
    tipo  = r.get("tipo","")
    web   = (r.get("web") or "").strip()

    if email not in VACIO and "@" in email:
        return ("IA - email", "IA",
                "Redacto email personalizado listo para enviar", email)
    if via == "formulario_web":
        return ("IA - formulario", "IA",
                "Preparo el texto de la solicitud para pegar en su formulario",
                web if web not in VACIO else "-")
    if via == "red_angels":
        return ("IA - formulario", "IA",
                "Preparo la solicitud de aplicacion a la red",
                web if web not in VACIO else "-")
    if via == "evento":
        return ("TU - evento", "TU",
                "Peticion de reunion en Fruit Attraction (6-8 oct) y cara a cara",
                web if web not in VACIO else "-")
    if via == "intro_calida":
        return ("TU - intro calida", "TU",
                "Solo funciona con presentacion de un contacto comun. Busca el puente",
                r.get("linkedin_url") if (r.get("linkedin_url") or "-") not in VACIO else "-")
    if tipo in ("industrial_agro","cadena_valor"):
        return ("TU - llamada", "TU",
                "Llamada directa a centralita y preguntar por compras o direccion",
                web if web not in VACIO else "-")
    return ("TU - llamada", "TU",
            "Sin canal publico: localizar telefono y llamar",
            web if web not in VACIO else "-")

VERDE="1F6F4A"; THIN=Side(style="thin", color="D5DDD8")
BORDER=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
FILL_IA=PatternFill("solid", fgColor="DCE9F7")   # azul: lo hace la IA
FILL_TU=PatternFill("solid", fgColor="FDF0D5")   # ambar: lo haces tu

def main():
    rows=list(csv.DictReader(open(SRC,encoding="utf-8-sig"),delimiter=";"))
    leads=[r for r in rows if puede_poner_capital(r)]
    for r in leads:
        c,q,acc,dato = canal(r)
        r["canal"],r["quien"],r["accion"],r["dato"] = c,q,acc,dato
        r["_s"]=int(float(r["score"] or 0))
    # Orden: primero lo que puede arrancar la IA (para lanzarlo ya), por score.
    leads.sort(key=lambda r:(0 if r["quien"]=="IA" else 1, -r["_s"], r["entidad"].lower()))

    COLS=[("orden",6),("entidad",34),("tipo",16),("ciudad",16),("score",7),("tier",6),
          ("canal_primer_contacto",18),("quien",7),("accion_concreta",44),
          ("dato_de_contacto",32),("ticket_min_eur",12),("ticket_max_eur",12),
          ("angulo_pitch",50),("red_flags",38),("web",30),("confianza",10),("estado",14),
          ("ESTADO_CONTACTO",18),("FECHA_CONTACTO",15),("RESPUESTA",14),("NDA",8),
          ("NUMEROS",10),("VISITA",9),("PROXIMA_ACCION",30),("FECHA_PROXIMA",15),("NOTAS",40)]
    nombres=[c for c,_ in COLS]

    wb=openpyxl.load_workbook(DST)
    if "Leads capital" in wb.sheetnames: del wb["Leads capital"]
    ws=wb.create_sheet("Leads capital", 1)   # justo despues de Portada

    ws.append(nombres)
    for i,(c,w) in enumerate(COLS, start=1):
        cell=ws.cell(row=1,column=i)
        cell.font=Font(bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=VERDE)
        cell.alignment=Alignment(vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=30

    for n,r in enumerate(leads, start=1):
        ws.append([n, r["entidad"], r["tipo"], r["ciudad"], r["_s"], r["tier"],
                   r["canal"], r["quien"], r["accion"], r["dato"],
                   r["ticket_min_eur"], r["ticket_max_eur"], r["angulo_pitch"],
                   r["red_flags"], r["web"], r["confianza"], r["estado"],
                   "pendiente","","","","","","","",""])
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row,max_col=len(nombres)):
        quien=ws.cell(row=row[0].row,column=8).value
        fill=FILL_IA if quien=="IA" else FILL_TU
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.border=BORDER; cell.font=Font(size=9); cell.fill=fill

    ws.freeze_panes="C2"
    ws.auto_filter.ref=f"A1:{get_column_letter(len(nombres))}{ws.max_row}"

    dv=DataValidation(type="list",formula1='"pendiente,contactado,sin respuesta,interesado,reunion agendada,NDA firmado,numeros enviados,visita hecha,oferta,descartado"',allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"R2:R{ws.max_row}")
    dv2=DataValidation(type="list",formula1='"si,no,pendiente"',allow_blank=True)
    ws.add_data_validation(dv2)
    for col in ("U","V","W"): dv2.add(f"{col}2:{col}{ws.max_row}")

    wb.save(DST)
    c=Counter(r["quien"] for r in leads); ch=Counter(r["canal"] for r in leads)
    print(f"Pestana 'Leads capital': {len(leads)} entidades que pueden aportar capital")
    print(f"  Puede arrancar la IA: {c['IA']}   ·   Requiere al fundador: {c['TU']}")
    for k,v in ch.most_common(): print(f"    {k:22s} {v}")
    print(f"\n  Tier A/B dentro de la lista: {sum(1 for r in leads if r['tier'] in ('A','B'))}")

if __name__=="__main__":
    main()
